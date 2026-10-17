from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
import secrets

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import delete, func, or_, select
from sqlalchemy.orm import Session, selectinload

from ..config import get_settings
from ..db import get_db
from ..deps import admin_user, owner_user, permission_required
from ..models import (
    AdminSecurity,
    AssistantQueryLog,
    AuditLog,
    AuthorizedAdminEmail,
    ChatKnowledge,
    Choice,
    ContentRevision,
    Notification,
    PageVisit,
    Question,
    Quiz,
    QuizAnswer,
    QuizAttempt,
    QuizPathConfig,
    RefreshToken,
    SitePage,
    SiteSetting,
    StudyMonth,
    StudentArchive,
    User,
    UserPresence,
)
from ..schemas import (
    AdminPermissionsIn,
    AuthorizedAdminIn,
    PageUpdateIn,
    QuizPathConfigIn,
    QuizUpdateIn,
    StudyMonthIn,
)
from ..security import hash_password

router = APIRouter(prefix="/admin", tags=["admin-plus"], dependencies=[Depends(admin_user)])

ALL_PERMISSIONS = [
    "dashboard", "students", "results", "reviews", "months", "quizzes", "content",
    "assistant", "settings", "analytics", "notifications", "admins", "security", "backup", "audit",
]
DEFAULT_ADMIN_PERMISSIONS = [x for x in ALL_PERMISSIONS if x not in {"security", "backup", "audit"}]


def _audit(db: Session, actor: int, action: str, entity: str, entity_id=None, metadata: dict | None = None):
    db.add(AuditLog(
        actor_user_id=actor,
        action=action,
        entity_type=entity,
        entity_id=str(entity_id) if entity_id is not None else None,
        metadata_json=json.dumps(metadata or {}, ensure_ascii=False, separators=(",", ":")),
    ))


def _parse_dt(value: str | datetime | None) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed
    except ValueError as exc:
        raise HTTPException(422, "التاريخ أو الوقت غير صحيح") from exc


def _permissions_for(db: Session, user: User) -> list[str]:
    if user.role == "owner":
        return ALL_PERMISSIONS
    sec = db.scalar(select(AdminSecurity).where(AdminSecurity.user_id == user.id))
    if not sec or not getattr(sec, "permissions_customized", False):
        return DEFAULT_ADMIN_PERMISSIONS
    try:
        values = json.loads(sec.permissions_json)
    except Exception:
        values = DEFAULT_ADMIN_PERMISSIONS
    return [x for x in values if x in ALL_PERMISSIONS]


@router.get("/permissions/me")
def my_permissions(db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    return {"role": actor.role, "permissions": _permissions_for(db, actor)}


@router.get("/authorized-admins", dependencies=[Depends(owner_user)])
def authorized_admins(db: Session = Depends(get_db)):
    rows = db.scalars(select(AuthorizedAdminEmail).order_by(AuthorizedAdminEmail.created_at)).all()
    users = {u.email: u for u in db.scalars(select(User).where(User.role.in_(["admin", "owner"]))).all()}
    return [{
        "id": row.id,
        "email": row.email,
        "role": row.role,
        "active": row.is_active,
        "user_id": users.get(row.email).id if users.get(row.email) else None,
        "account_active": users.get(row.email).is_active if users.get(row.email) else None,
        "permissions": _permissions_for(db, users.get(row.email)) if users.get(row.email) else [],
    } for row in rows]


@router.post("/authorized-admins", status_code=201, dependencies=[Depends(owner_user)])
def add_authorized_admin(data: AuthorizedAdminIn, db: Session = Depends(get_db), actor: User = Depends(owner_user)):
    email = str(data.email).lower().strip()
    existing = db.scalar(select(AuthorizedAdminEmail).where(AuthorizedAdminEmail.email == email))
    if existing:
        existing.role = data.role
        existing.is_active = True
        existing.created_by = actor.id
        row = existing
    else:
        row = AuthorizedAdminEmail(email=email, role=data.role, created_by=actor.id)
        db.add(row)
    user = db.scalar(select(User).where(User.email == email))
    if user and user.role in {"admin", "owner"}:
        if user.id != actor.id:
            user.role = data.role
        user.is_active = True
    _audit(db, actor.id, "authorized_admin_saved", "authorized_admin", email, {"role": data.role})
    db.commit()
    return {"ok": True, "email": email, "role": data.role}


@router.delete("/authorized-admins/{admin_id}", dependencies=[Depends(owner_user)])
def remove_authorized_admin(admin_id: int, db: Session = Depends(get_db), actor: User = Depends(owner_user)):
    row = db.get(AuthorizedAdminEmail, admin_id)
    if not row:
        raise HTTPException(404, "الأدمن غير موجود")
    if row.email == actor.email:
        raise HTTPException(409, "مينفعش توقف إيميل حسابك الحالي")
    row.is_active = False
    target = db.scalar(select(User).where(User.email == row.email))
    if target and target.role in {"admin", "owner"}:
        target.is_active = False
        now = datetime.now(timezone.utc)
        db.query(RefreshToken).filter(RefreshToken.user_id == target.id, RefreshToken.revoked_at.is_(None)).update({"revoked_at": now})
    _audit(db, actor.id, "authorized_admin_disabled", "authorized_admin", row.email)
    db.commit()
    return {"ok": True}


@router.put("/admins/{user_id}/permissions", dependencies=[Depends(owner_user)])
def update_permissions(user_id: int, data: AdminPermissionsIn, db: Session = Depends(get_db), actor: User = Depends(owner_user)):
    user = db.get(User, user_id)
    if not user or user.role not in {"admin", "owner"}:
        raise HTTPException(404, "حساب الأدمن غير موجود")
    if user.role == "owner":
        raise HTTPException(409, "المالك عنده كل الصلاحيات تلقائيًا")
    clean = sorted(set(x for x in data.permissions if x in ALL_PERMISSIONS))
    sec = db.scalar(select(AdminSecurity).where(AdminSecurity.user_id == user.id))
    if not sec:
        sec = AdminSecurity(user_id=user.id)
        db.add(sec)
    sec.permissions_json = json.dumps(clean, ensure_ascii=False)
    sec.permissions_customized = True
    _audit(db, actor.id, "admin_permissions_changed", "user", user.id, {"permissions": clean})
    db.commit()
    return {"permissions": clean}


@router.get("/months", dependencies=[Depends(permission_required("months"))])
def months(db: Session = Depends(get_db)):
    rows = db.scalars(select(StudyMonth).order_by(StudyMonth.sort_order, StudyMonth.id)).all()
    counts = dict(db.execute(select(QuizPathConfig.month_id, func.count()).group_by(QuizPathConfig.month_id)).all())
    return [{
        "id": m.id, "title": m.title, "stage": m.stage, "education_type": m.education_type,
        "grade": m.grade, "sort_order": m.sort_order, "is_visible": m.is_visible,
        "starts_at": m.starts_at, "ends_at": m.ends_at, "quiz_count": counts.get(m.id, 0),
    } for m in rows]


@router.post("/months", status_code=201, dependencies=[Depends(permission_required("months"))])
def create_month(data: StudyMonthIn, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    row = StudyMonth(
        title=data.title.strip(), stage=data.stage, education_type=data.education_type, grade=data.grade,
        sort_order=data.sort_order, is_visible=data.is_visible,
        starts_at=_parse_dt(data.starts_at), ends_at=_parse_dt(data.ends_at),
    )
    db.add(row); db.flush(); _audit(db, actor.id, "month_created", "study_month", row.id); db.commit(); db.refresh(row)
    return {"id": row.id}


@router.put("/months/{month_id}", dependencies=[Depends(permission_required("months"))])
def update_month(month_id: int, data: StudyMonthIn, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    row = db.get(StudyMonth, month_id)
    if not row:
        raise HTTPException(404, "الشهر غير موجود")
    row.title = data.title.strip(); row.stage = data.stage; row.education_type = data.education_type; row.grade = data.grade
    row.sort_order = data.sort_order; row.is_visible = data.is_visible
    row.starts_at = _parse_dt(data.starts_at); row.ends_at = _parse_dt(data.ends_at)
    _audit(db, actor.id, "month_updated", "study_month", row.id); db.commit()
    return {"ok": True}


@router.delete("/months/{month_id}", dependencies=[Depends(permission_required("months"))])
def delete_month(month_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    row = db.get(StudyMonth, month_id)
    if not row:
        raise HTTPException(404, "الشهر غير موجود")
    for cfg in db.scalars(select(QuizPathConfig).where(QuizPathConfig.month_id == month_id)).all():
        cfg.month_id = None
    db.delete(row); _audit(db, actor.id, "month_deleted", "study_month", month_id); db.commit()
    return {"ok": True}


@router.get("/quiz-path", dependencies=[Depends(permission_required("months"))])
def quiz_path_admin(db: Session = Depends(get_db)):
    configs = db.scalars(select(QuizPathConfig).order_by(QuizPathConfig.month_id, QuizPathConfig.sort_order)).all()
    quizzes = {q.id: q for q in db.scalars(select(Quiz)).all()}
    return [{
        "quiz_id": c.quiz_id,
        "quiz": quizzes[c.quiz_id].title_ar if c.quiz_id in quizzes else f"Quiz #{c.quiz_id}",
        "month_id": c.month_id,
        "week_label": c.week_label,
        "sort_order": c.sort_order,
        "requires_previous": c.requires_previous,
        "unlock_after_quiz_id": c.unlock_after_quiz_id,
        "min_previous_score": c.min_previous_score,
        "show_when_locked": c.show_when_locked,
    } for c in configs]


@router.put("/quiz-path/{quiz_id}", dependencies=[Depends(permission_required("months"))])
def save_quiz_path(quiz_id: int, data: QuizPathConfigIn, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    quiz = db.get(Quiz, quiz_id)
    if not quiz:
        raise HTTPException(404, "الاختبار غير موجود")
    if data.month_id and not db.get(StudyMonth, data.month_id):
        raise HTTPException(404, "الشهر غير موجود")
    if data.unlock_after_quiz_id == quiz_id:
        raise HTTPException(422, "الاختبار لا يمكن أن يعتمد على نفسه")
    row = db.scalar(select(QuizPathConfig).where(QuizPathConfig.quiz_id == quiz_id))
    if not row:
        row = QuizPathConfig(quiz_id=quiz_id)
        db.add(row)
    for field, value in data.model_dump().items():
        setattr(row, field, value)
    _audit(db, actor.id, "quiz_path_saved", "quiz", quiz_id, data.model_dump())
    db.commit()
    return {"ok": True}


@router.delete("/quiz-path/{quiz_id}", dependencies=[Depends(permission_required("months"))])
def delete_quiz_path(quiz_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    row = db.scalar(select(QuizPathConfig).where(QuizPathConfig.quiz_id == quiz_id))
    if row:
        db.delete(row); _audit(db, actor.id, "quiz_path_removed", "quiz", quiz_id); db.commit()
    return {"ok": True}


@router.put("/quizzes/{quiz_id}/details", dependencies=[Depends(permission_required("quizzes"))])
def update_quiz_details(quiz_id: int, data: QuizUpdateIn, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    quiz = db.get(Quiz, quiz_id)
    if not quiz:
        raise HTTPException(404, "الاختبار غير موجود")
    payload = data.model_dump(exclude_unset=True)
    before = {}
    for key in payload:
        old_value = getattr(quiz, key, None)
        before[key] = old_value.isoformat() if isinstance(old_value, datetime) else old_value
    for key in ("starts_at", "ends_at"):
        if key in payload:
            payload[key] = _parse_dt(payload[key])
    for key, value in payload.items():
        setattr(quiz, key, value)
    after = {key: (value.isoformat() if isinstance(value, datetime) else value) for key, value in payload.items()}
    _audit(db, actor.id, "quiz_updated", "quiz", quiz.id, {"before": before, "after": after})
    db.commit()
    return {"ok": True}


def _save_revision(db: Session, page: SitePage, actor_id: int):
    db.add(ContentRevision(
        page_id=page.id, title=page.title, path=page.path, is_visible=page.is_visible,
        sort_order=page.sort_order, content_json=page.content_json, changed_by=actor_id,
    ))


@router.put("/pages/{page_id}/content", dependencies=[Depends(permission_required("content"))])
def update_page_content(page_id: int, data: PageUpdateIn, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    page = db.get(SitePage, page_id)
    if not page or page.archived_at is not None:
        raise HTTPException(404, "الصفحة غير موجودة")
    before = {"title": page.title, "path": page.path, "is_visible": page.is_visible, "sort_order": page.sort_order, "content_json": page.content_json}
    _save_revision(db, page, actor.id)
    payload = data.model_dump(exclude_unset=True, exclude_none=True)
    if "content_json" in payload:
        try:
            json.loads(payload["content_json"])
        except Exception as exc:
            raise HTTPException(422, "محتوى الصفحة لازم يكون JSON صحيح") from exc
    for key, value in payload.items():
        setattr(page, key, value)
    after = {key: getattr(page, key) for key in payload}
    _audit(db, actor.id, "page_content_updated", "site_page", page.id, {"before": {k: before[k] for k in payload}, "after": after})
    db.commit()
    return {"ok": True}


@router.get("/pages/{page_id}/revisions", dependencies=[Depends(permission_required("content"))])
def page_revisions(page_id: int, db: Session = Depends(get_db)):
    rows = db.scalars(select(ContentRevision).where(ContentRevision.page_id == page_id).order_by(ContentRevision.created_at.desc()).limit(30)).all()
    return [{"id": r.id, "title": r.title, "created_at": r.created_at, "changed_by": r.changed_by} for r in rows]


@router.post("/pages/{page_id}/rollback/{revision_id}", dependencies=[Depends(permission_required("content"))])
def rollback_page(page_id: int, revision_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    page = db.get(SitePage, page_id); revision = db.get(ContentRevision, revision_id)
    if not page or not revision or revision.page_id != page_id:
        raise HTTPException(404, "النسخة السابقة غير موجودة")
    _save_revision(db, page, actor.id)
    page.title = revision.title; page.path = revision.path; page.is_visible = revision.is_visible
    page.sort_order = revision.sort_order; page.content_json = revision.content_json
    _audit(db, actor.id, "page_rolled_back", "site_page", page.id, {"revision_id": revision.id})
    db.commit()
    return {"ok": True}


@router.get("/analytics", dependencies=[Depends(permission_required("analytics"))])
def analytics(db: Session = Depends(get_db)):
    total = db.scalar(select(func.count()).select_from(PageVisit)) or 0
    by_page_rows = db.execute(select(PageVisit.path, func.count()).group_by(PageVisit.path).order_by(func.count().desc()).limit(50)).all()
    cutoff = datetime.now(timezone.utc) - timedelta(days=7)
    day_expr = func.date(PageVisit.visited_at)
    day_rows = db.execute(select(day_expr, func.count()).where(PageVisit.visited_at >= cutoff).group_by(day_expr).order_by(day_expr)).all()
    online_cutoff = datetime.now(timezone.utc) - timedelta(minutes=2)
    online = db.scalar(select(func.count()).select_from(UserPresence).where(UserPresence.last_seen >= online_cutoff)) or 0
    grade_rows = db.execute(select(User.grade, func.count()).where(User.role == "student").group_by(User.grade)).all()
    return {
        "total": total,
        "by_page": [{"path": p, "count": c} for p, c in by_page_rows],
        "last_7_days": [{"date": str(d), "count": c} for d, c in day_rows],
        "online_now": online,
        "students_by_grade": [{"grade": g or "غير محدد", "count": c} for g, c in grade_rows],
    }


@router.get("/presence", dependencies=[Depends(permission_required("analytics"))])
def presence(db: Session = Depends(get_db)):
    rows = db.execute(
        select(UserPresence, User).join(User, User.id == UserPresence.user_id).order_by(UserPresence.last_seen.desc()).limit(300)
    ).all()
    now = datetime.now(timezone.utc)
    return [{
        "user_id": u.id, "name": f"{u.first_name} {u.last_name}", "email": u.email,
        "role": u.role, "path": p.path, "last_seen": p.last_seen,
        "online": (now - (p.last_seen if p.last_seen.tzinfo else p.last_seen.replace(tzinfo=timezone.utc))) <= timedelta(minutes=2),
    } for p, u in rows]


@router.get("/assistant/questions", dependencies=[Depends(permission_required("assistant"))])
def assistant_questions(db: Session = Depends(get_db)):
    rows = db.execute(
        select(AssistantQueryLog.question, AssistantQueryLog.source, func.count())
        .group_by(AssistantQueryLog.question, AssistantQueryLog.source)
        .order_by(func.count().desc()).limit(100)
    ).all()
    return [{"question": q, "source": s, "count": c} for q, s, c in rows]


@router.get("/question-analytics", dependencies=[Depends(permission_required("analytics"))])
def question_analytics(db: Session = Depends(get_db)):
    # Portable aggregation: compute over a bounded result set instead of DB-specific boolean casts.
    answers = db.execute(
        select(QuizAnswer.is_correct, Question.id, Question.prompt, Quiz.title_ar)
        .join(Question, Question.id == QuizAnswer.question_id)
        .join(Quiz, Quiz.id == Question.quiz_id)
        .where(QuizAnswer.is_correct.is_not(None))
        .limit(20000)
    ).all()
    stats: dict[int, dict] = {}
    for is_correct, qid, prompt, quiz_title in answers:
        item = stats.setdefault(qid, {"question_id": qid, "question": prompt, "quiz": quiz_title, "answers": 0, "correct": 0})
        item["answers"] += 1
        if is_correct:
            item["correct"] += 1
    result = []
    for item in stats.values():
        item["accuracy"] = round(item["correct"] / item["answers"] * 100, 1) if item["answers"] else None
        result.append(item)
    result.sort(key=lambda x: (x["accuracy"] if x["accuracy"] is not None else 101, -x["answers"]))
    return result[:100]


@router.get("/search")
def global_search(q: str = Query(min_length=2, max_length=100), db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    term = f"%{q.strip()}%"
    permissions = set(_permissions_for(db, actor))
    users = db.scalars(select(User).where(User.role == "student", or_(User.first_name.ilike(term), User.last_name.ilike(term), User.email.ilike(term), User.phone.ilike(term))).limit(10)).all() if "students" in permissions else []
    quizzes = db.scalars(select(Quiz).where(Quiz.title_ar.ilike(term), Quiz.is_archived.is_(False)).limit(10)).all() if "quizzes" in permissions else []
    pages = db.scalars(select(SitePage).where(SitePage.title.ilike(term), SitePage.archived_at.is_(None)).limit(10)).all() if "content" in permissions else []
    return {
        "users": [{"id": u.id, "title": f"{u.first_name} {u.last_name}", "subtitle": u.phone} for u in users],
        "quizzes": [{"id": x.id, "title": x.title_ar, "subtitle": x.slug} for x in quizzes],
        "pages": [{"id": x.id, "title": x.title, "subtitle": x.path} for x in pages],
    }


@router.post("/users/{user_id}/archive", dependencies=[Depends(permission_required("students"))])
def archive_student(user_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    user = db.get(User, user_id)
    if not user or user.role != "student":
        raise HTTPException(404, "حساب الطالب غير موجود")
    existing = db.get(StudentArchive, user.id)
    if not existing:
        db.add(StudentArchive(user_id=user.id, archived_by=actor.id))
    user.is_active = False
    now = datetime.now(timezone.utc)
    db.query(RefreshToken).filter(RefreshToken.user_id == user.id, RefreshToken.revoked_at.is_(None)).update({"revoked_at": now})
    _audit(db, actor.id, "student_archived", "user", user.id, {"email": user.email})
    db.commit()
    return {"ok": True}


@router.get("/trash", dependencies=[Depends(owner_user)])
def trash(db: Session = Depends(get_db)):
    quizzes = db.scalars(select(Quiz).where(Quiz.is_archived.is_(True)).order_by(Quiz.created_at.desc())).all()
    pages = db.scalars(select(SitePage).where(SitePage.archived_at.is_not(None)).order_by(SitePage.archived_at.desc())).all()
    archived_students = db.execute(
        select(StudentArchive, User).join(User, User.id == StudentArchive.user_id).order_by(StudentArchive.archived_at.desc())
    ).all()
    return {
        "quizzes": [{"id": q.id, "title": q.title_ar, "archived_at": q.created_at} for q in quizzes],
        "pages": [{"id": p.id, "title": p.title, "archived_at": p.archived_at} for p in pages],
        "students": [{"id": u.id, "title": f"{u.first_name} {u.last_name}", "email": u.email, "archived_at": a.archived_at} for a, u in archived_students],
    }


@router.post("/trash/quizzes/{quiz_id}/restore", dependencies=[Depends(owner_user)])
def restore_quiz(quiz_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    q = db.get(Quiz, quiz_id)
    if not q:
        raise HTTPException(404, "الاختبار غير موجود")
    q.is_archived = False; _audit(db, actor.id, "quiz_restored", "quiz", q.id); db.commit(); return {"ok": True}


@router.post("/trash/pages/{page_id}/restore", dependencies=[Depends(owner_user)])
def restore_page(page_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    p = db.get(SitePage, page_id)
    if not p:
        raise HTTPException(404, "الصفحة غير موجودة")
    p.archived_at = None; _audit(db, actor.id, "page_restored", "site_page", p.id); db.commit(); return {"ok": True}


@router.post("/trash/students/{user_id}/restore", dependencies=[Depends(owner_user)])
def restore_student(user_id: int, db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    user = db.get(User, user_id)
    archive = db.get(StudentArchive, user_id)
    if not user or user.role != "student" or not archive:
        raise HTTPException(404, "الطالب غير موجود في المحذوفات")
    db.delete(archive)
    user.is_active = True
    _audit(db, actor.id, "student_restored", "user", user.id, {"email": user.email})
    db.commit()
    return {"ok": True}


@router.get("/system-status", dependencies=[Depends(owner_user)])
def system_status(db: Session = Depends(get_db)):
    settings = get_settings()
    db.execute(select(func.count()).select_from(User)).scalar_one()
    return {
        "backend": "ok",
        "database": "ok",
        "gemini_configured": bool(settings.gemini_api_key.strip()),
        "gemini_model": settings.gemini_model,
        "email_configured": bool(settings.smtp_host and settings.smtp_username and settings.smtp_password),
        "environment": settings.env,
    }


@router.get("/sessions", dependencies=[Depends(owner_user)])
def sessions(db: Session = Depends(get_db)):
    now = datetime.now(timezone.utc)
    rows = db.execute(
        select(User.id, User.first_name, User.last_name, User.email, func.count(RefreshToken.id), func.max(RefreshToken.expires_at))
        .join(RefreshToken, RefreshToken.user_id == User.id)
        .where(RefreshToken.revoked_at.is_(None), RefreshToken.expires_at > now)
        .group_by(User.id, User.first_name, User.last_name, User.email)
    ).all()
    return [{"user_id": uid, "name": f"{first} {last}", "email": email, "active_sessions": count, "latest_expiry": expiry} for uid, first, last, email, count, expiry in rows]


@router.post("/sessions/{user_id}/revoke", dependencies=[Depends(owner_user)])
def revoke_sessions(user_id: int, db: Session = Depends(get_db), actor: User = Depends(owner_user)):
    now = datetime.now(timezone.utc)
    count = db.query(RefreshToken).filter(RefreshToken.user_id == user_id, RefreshToken.revoked_at.is_(None)).update({"revoked_at": now})
    _audit(db, actor.id, "sessions_revoked", "user", user_id, {"count": count}); db.commit(); return {"revoked": count}


@router.get("/exports/users.xlsx", dependencies=[Depends(permission_required("students"))])
def export_users(db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    rows = db.scalars(select(User).where(User.role == "student").order_by(User.created_at)).all()
    wb = Workbook(); ws = wb.active; ws.title = "الطلاب"; ws.sheet_view.rightToLeft = True
    headers = ["الاسم الأول", "الاسم الأخير", "الإيميل", "الهاتف", "ولي الأمر", "المرحلة", "نوع التعليم", "الصف", "المدرسة", "نشط", "تاريخ التسجيل"]
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(horizontal="center")
    for u in rows:
        ws.append([u.first_name, u.last_name, u.email, u.phone, u.guardian_phone or "", u.stage or "", u.education_type or "", u.grade or "", u.school_name or "", "نعم" if u.is_active else "لا", u.created_at.isoformat()])
    stream = BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="raed-students.xlsx"'})


@router.get("/imports/users-template.xlsx", dependencies=[Depends(permission_required("students"))])
def import_template():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "استيراد الطلاب"; ws.sheet_view.rightToLeft = True
    ws.append(["first_name", "last_name", "email", "phone", "guardian_phone", "stage", "education_type", "grade", "school_name", "password"])
    ws.append(["أحمد", "محمد", "student@example.com", "01000000000", "01011111111", "ثانوي", "عام", "الصف الأول الثانوي", "مدرسة مثال", "StrongPass1"])
    stream = BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="students-import-template.xlsx"'})


@router.post("/imports/users.xlsx", dependencies=[Depends(permission_required("students"))])
def import_users(file: UploadFile = File(...), db: Session = Depends(get_db), actor: User = Depends(admin_user)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "ارفع ملف Excel بصيغة xlsx")
    from openpyxl import load_workbook
    content = file.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "ملف Excel أكبر من 5MB")
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True); ws = wb.active
    except Exception as exc:
        raise HTTPException(422, "ملف Excel غير صالح") from exc
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(422, "الملف فارغ")
    headers = [str(x or "").strip() for x in rows[0]]
    required = {"first_name", "last_name", "email", "phone"}
    if not required.issubset(headers):
        raise HTTPException(422, "قالب الملف غير صحيح. نزّل قالب الاستيراد من لوحة الإدارة")
    created, skipped, temp_passwords = 0, [], []
    for index, values in enumerate(rows[1:], start=2):
        data = dict(zip(headers, values))
        email = str(data.get("email") or "").lower().strip(); phone = str(data.get("phone") or "").strip()
        if not email or not phone or db.scalar(select(User).where(or_(User.email == email, User.phone == phone))):
            skipped.append(index); continue
        password = str(data.get("password") or "").strip()
        if len(password) < 8 or not any(c.isupper() for c in password) or not any(c.islower() for c in password) or not any(c.isdigit() for c in password):
            password = f"WelcomeA1{secrets.randbelow(9000)+1000}"
            temp_passwords.append({"email": email, "password": password})
        user = User(
            first_name=str(data.get("first_name") or "").strip()[:80], last_name=str(data.get("last_name") or "").strip()[:80],
            email=email, phone=phone, guardian_phone=str(data.get("guardian_phone") or "").strip() or None,
            stage=str(data.get("stage") or "").strip() or None, education_type=str(data.get("education_type") or "").strip() or None,
            grade=str(data.get("grade") or "").strip() or None, school_name=str(data.get("school_name") or "").strip() or None,
            password_hash=hash_password(password), role="student", is_active=True,
        )
        db.add(user); created += 1
    _audit(db, actor.id, "students_imported", "user", None, {"created": created, "skipped": len(skipped)}); db.commit()
    return {"created": created, "skipped_rows": skipped, "temporary_passwords": temp_passwords}


def _snapshot(db: Session) -> dict:
    pages = db.scalars(select(SitePage)).all()
    months = db.scalars(select(StudyMonth)).all()
    paths = db.scalars(select(QuizPathConfig)).all()
    quizzes = db.scalars(select(Quiz).options(selectinload(Quiz.questions).selectinload(Question.choices))).all()
    students = db.scalars(select(User).where(User.role == "student").order_by(User.id)).all()
    archived_ids = {x.user_id for x in db.scalars(select(StudentArchive)).all()}
    attempts = db.scalars(
        select(QuizAttempt)
        .options(selectinload(QuizAttempt.user), selectinload(QuizAttempt.quiz), selectinload(QuizAttempt.answers))
        .order_by(QuizAttempt.id)
    ).all()
    question_ids = {a.question_id for attempt in attempts for a in attempt.answers}
    choice_ids = {a.choice_id for attempt in attempts for a in attempt.answers if a.choice_id}
    question_map = {q.id: q for q in db.scalars(select(Question).where(Question.id.in_(question_ids))).all()} if question_ids else {}
    choice_map = {c.id: c for c in db.scalars(select(Choice).where(Choice.id.in_(choice_ids))).all()} if choice_ids else {}
    return {
        "version": 3,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "site_settings": {x.key: x.value for x in db.scalars(select(SiteSetting)).all()},
        "pages": [{
            "backup_id": p.id, "slug": p.slug, "title": p.title, "path": p.path,
            "is_visible": p.is_visible, "sort_order": p.sort_order, "is_system": p.is_system,
            "content_json": p.content_json, "archived_at": p.archived_at.isoformat() if p.archived_at else None,
        } for p in pages],
        "months": [{
            "backup_id": m.id, "title": m.title, "stage": m.stage, "education_type": m.education_type,
            "grade": m.grade, "sort_order": m.sort_order, "is_visible": m.is_visible,
            "starts_at": m.starts_at.isoformat() if m.starts_at else None,
            "ends_at": m.ends_at.isoformat() if m.ends_at else None,
        } for m in months],
        "quizzes": [{
            "backup_id": q.id, "slug": q.slug, "title_ar": q.title_ar, "title_en": q.title_en,
            "description_ar": q.description_ar, "description_en": q.description_en,
            "stage": q.stage, "education_type": q.education_type, "grade": q.grade,
            "duration_minutes": q.duration_minutes, "passing_score": q.passing_score,
            "max_attempts": q.max_attempts, "randomize_questions": q.randomize_questions,
            "show_result_immediately": q.show_result_immediately,
            "starts_at": q.starts_at.isoformat() if q.starts_at else None,
            "ends_at": q.ends_at.isoformat() if q.ends_at else None,
            "is_published": q.is_published, "is_placement": q.is_placement, "is_archived": q.is_archived,
            "questions": [{
                "prompt": x.prompt, "explanation": x.explanation, "level": x.level,
                "question_type": x.question_type, "expected_answer": x.expected_answer,
                "position": x.position, "points": x.points, "is_active": x.is_active,
                "choices": [{"text": c.text, "is_correct": c.is_correct, "position": c.position} for c in x.choices],
            } for x in q.questions],
        } for q in quizzes],
        "quiz_path": [{
            "quiz_backup_id": x.quiz_id, "month_backup_id": x.month_id, "week_label": x.week_label,
            "sort_order": x.sort_order, "requires_previous": x.requires_previous,
            "unlock_after_quiz_backup_id": x.unlock_after_quiz_id,
            "min_previous_score": x.min_previous_score, "show_when_locked": x.show_when_locked,
        } for x in paths],
        "knowledge": [{
            "category": x.category, "question": x.question, "answer": x.answer,
            "keywords": x.keywords, "is_active": x.is_active,
        } for x in db.scalars(select(ChatKnowledge)).all()],
        # Password hashes, refresh tokens, OTP codes, API keys and admin secrets are deliberately excluded.
        "students": [{
            "backup_id": u.id, "first_name": u.first_name, "last_name": u.last_name,
            "email": u.email, "phone": u.phone, "guardian_phone": u.guardian_phone,
            "stage": u.stage, "education_type": u.education_type, "grade": u.grade,
            "school_name": u.school_name, "profile_image": u.profile_image,
            "is_active": u.is_active, "is_verified": u.is_verified,
            "archived": u.id in archived_ids,
            "created_at": u.created_at.isoformat() if u.created_at else None,
        } for u in students],
        "attempts": [{
            "backup_id": attempt.id,
            "student_email": attempt.user.email,
            "quiz_slug": attempt.quiz.slug,
            "started_at": attempt.started_at.isoformat() if attempt.started_at else None,
            "submitted_at": attempt.submitted_at.isoformat() if attempt.submitted_at else None,
            "score": attempt.score, "level_result": attempt.level_result,
            "needs_manual_review": attempt.needs_manual_review,
            "answers": [{
                "question_position": question_map[a.question_id].position if a.question_id in question_map else None,
                "choice_position": choice_map[a.choice_id].position if a.choice_id in choice_map else None,
                "text_answer": a.text_answer, "is_correct": a.is_correct, "awarded_points": a.awarded_points,
            } for a in attempt.answers],
        } for attempt in attempts if attempt.user and attempt.user.role == "student" and attempt.quiz],
    }


@router.get("/backup/export.json", dependencies=[Depends(owner_user)])
def backup_export(db: Session = Depends(get_db)):
    data = json.dumps(_snapshot(db), ensure_ascii=False, indent=2).encode("utf-8")
    return StreamingResponse(BytesIO(data), media_type="application/json; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="raed-full-safe-backup.json"'})


@router.post("/backup/restore", dependencies=[Depends(owner_user)])
def backup_restore(file: UploadFile = File(...), db: Session = Depends(get_db), actor: User = Depends(owner_user)):
    raw = file.file.read(50 * 1024 * 1024 + 1)
    if len(raw) > 50 * 1024 * 1024:
        raise HTTPException(413, "ملف النسخة الاحتياطية كبير جدًا")
    try:
        data = json.loads(raw.decode("utf-8"))
    except Exception as exc:
        raise HTTPException(422, "ملف النسخة الاحتياطية غير صالح") from exc
    if data.get("version") not in {1, 2, 3}:
        raise HTTPException(422, "إصدار النسخة الاحتياطية غير مدعوم")

    # Settings and editable site content are safe to upsert because they do not own student/result data.
    for key, value in (data.get("site_settings") or {}).items():
        row = db.get(SiteSetting, key)
        if row:
            row.value = str(value)
        else:
            db.add(SiteSetting(key=str(key)[:120], value=str(value)))

    for item in data.get("pages") or []:
        slug = str(item.get("slug") or "").strip()
        if not slug:
            continue
        row = db.scalar(select(SitePage).where(SitePage.slug == slug))
        if not row:
            row = SitePage(slug=slug, title=str(item.get("title") or slug), path=str(item.get("path") or f"/p/{slug}"))
            db.add(row)
            db.flush()
        row.title = str(item.get("title") or row.title)
        row.path = str(item.get("path") or row.path)
        row.is_visible = bool(item.get("is_visible", True))
        row.sort_order = int(item.get("sort_order", 0))
        row.content_json = str(item.get("content_json") or "{}")
        if data.get("version") in {2, 3}:
            row.archived_at = _parse_dt(item.get("archived_at"))

    month_id_map: dict[int, int] = {}
    for item in data.get("months") or []:
        title = str(item.get("title") or "").strip()
        if not title:
            continue
        row = db.scalar(select(StudyMonth).where(
            StudyMonth.title == title,
            StudyMonth.stage == item.get("stage"),
            StudyMonth.education_type == item.get("education_type"),
            StudyMonth.grade == item.get("grade"),
        ))
        if not row:
            row = StudyMonth(title=title)
            db.add(row)
            db.flush()
        row.stage = item.get("stage") or None
        row.education_type = item.get("education_type") or None
        row.grade = item.get("grade") or None
        row.sort_order = int(item.get("sort_order", 0))
        row.is_visible = bool(item.get("is_visible", True))
        row.starts_at = _parse_dt(item.get("starts_at"))
        row.ends_at = _parse_dt(item.get("ends_at"))
        if item.get("backup_id") is not None:
            month_id_map[int(item["backup_id"])] = row.id

    quiz_id_map: dict[int, int] = {}
    restored_question_banks = 0
    for item in data.get("quizzes") or []:
        slug = str(item.get("slug") or "").strip()
        if not slug:
            continue
        row = db.scalar(select(Quiz).where(Quiz.slug == slug).options(selectinload(Quiz.questions).selectinload(Question.choices)))
        is_new = row is None
        if not row:
            row = Quiz(slug=slug, title_ar=str(item.get("title_ar") or slug), title_en=str(item.get("title_en") or ""))
            db.add(row)
            db.flush()
        row.title_ar = str(item.get("title_ar") or row.title_ar)
        row.title_en = str(item.get("title_en") or "")
        row.description_ar = str(item.get("description_ar") or "")
        row.description_en = str(item.get("description_en") or "")
        row.stage = item.get("stage") or None
        row.education_type = item.get("education_type") or None
        row.grade = item.get("grade") or None
        row.duration_minutes = int(item.get("duration_minutes", 30))
        row.passing_score = int(item.get("passing_score", 60))
        row.max_attempts = int(item.get("max_attempts", 2))
        row.randomize_questions = bool(item.get("randomize_questions", False))
        row.show_result_immediately = bool(item.get("show_result_immediately", True))
        row.starts_at = _parse_dt(item.get("starts_at"))
        row.ends_at = _parse_dt(item.get("ends_at"))
        row.is_published = bool(item.get("is_published", False))
        row.is_placement = bool(item.get("is_placement", False))
        row.is_archived = bool(item.get("is_archived", False))
        db.flush()
        if item.get("backup_id") is not None:
            quiz_id_map[int(item["backup_id"])] = row.id

        has_attempts = db.scalar(select(QuizAttempt.id).where(QuizAttempt.quiz_id == row.id).limit(1)) is not None
        # Question IDs are referenced by historic answers. Only replace the bank when it is safe.
        if is_new or not has_attempts:
            for question in list(row.questions):
                db.delete(question)
            db.flush()
            for q_item in item.get("questions") or []:
                question = Question(
                    quiz_id=row.id,
                    prompt=str(q_item.get("prompt") or ""),
                    explanation=str(q_item.get("explanation") or ""),
                    level=str(q_item.get("level") or "A1")[:10],
                    question_type=str(q_item.get("question_type") or "mcq"),
                    expected_answer=q_item.get("expected_answer"),
                    position=int(q_item.get("position", 1)),
                    points=int(q_item.get("points", 1)),
                    is_active=bool(q_item.get("is_active", True)),
                )
                db.add(question)
                db.flush()
                for c_item in q_item.get("choices") or []:
                    db.add(Choice(
                        question_id=question.id,
                        text=str(c_item.get("text") or ""),
                        is_correct=bool(c_item.get("is_correct", False)),
                        position=int(c_item.get("position", 1)),
                    ))
            restored_question_banks += 1

    if data.get("version") in {2, 3}:
        for item in data.get("quiz_path") or []:
            source_quiz_id = item.get("quiz_backup_id")
            target_quiz_id = quiz_id_map.get(int(source_quiz_id)) if source_quiz_id is not None else None
            if not target_quiz_id:
                continue
            row = db.scalar(select(QuizPathConfig).where(QuizPathConfig.quiz_id == target_quiz_id))
            if not row:
                row = QuizPathConfig(quiz_id=target_quiz_id)
                db.add(row)
            source_month_id = item.get("month_backup_id")
            source_unlock_id = item.get("unlock_after_quiz_backup_id")
            row.month_id = month_id_map.get(int(source_month_id)) if source_month_id is not None else None
            row.week_label = item.get("week_label") or None
            row.sort_order = int(item.get("sort_order", 0))
            row.requires_previous = bool(item.get("requires_previous", False))
            row.unlock_after_quiz_id = quiz_id_map.get(int(source_unlock_id)) if source_unlock_id is not None else None
            row.min_previous_score = int(item.get("min_previous_score", 60))
            row.show_when_locked = bool(item.get("show_when_locked", True))

    for item in data.get("knowledge") or []:
        question = str(item.get("question") or "").strip()
        if not question:
            continue
        row = db.scalar(select(ChatKnowledge).where(ChatKnowledge.question == question))
        if not row:
            row = ChatKnowledge(category=str(item.get("category") or "عام"), question=question, answer="")
            db.add(row)
        row.category = str(item.get("category") or row.category)
        row.answer = str(item.get("answer") or "")
        row.keywords = str(item.get("keywords") or "")
        row.is_active = bool(item.get("is_active", True))

    restored_students = 0
    restored_attempts = 0
    if data.get("version") == 3:
        for item in data.get("students") or []:
            email = str(item.get("email") or "").lower().strip()
            phone = str(item.get("phone") or "").strip()
            if not email or not phone:
                continue
            user = db.scalar(select(User).where(User.email == email))
            if not user:
                phone_owner = db.scalar(select(User).where(User.phone == phone))
                if phone_owner:
                    # Do not overwrite a different account that already owns the phone number.
                    continue
                user = User(
                    first_name=str(item.get("first_name") or "طالب")[:80],
                    last_name=str(item.get("last_name") or "مستر رائد")[:80],
                    email=email,
                    phone=phone,
                    password_hash=hash_password(secrets.token_urlsafe(32)),
                    role="student",
                )
                db.add(user); db.flush(); restored_students += 1
            if user.role != "student":
                continue
            user.first_name = str(item.get("first_name") or user.first_name)[:80]
            user.last_name = str(item.get("last_name") or user.last_name)[:80]
            user.guardian_phone = item.get("guardian_phone") or None
            user.stage = item.get("stage") or None
            user.education_type = item.get("education_type") or None
            user.grade = item.get("grade") or None
            user.school_name = item.get("school_name") or None
            user.profile_image = item.get("profile_image") or None
            user.is_verified = bool(item.get("is_verified", False))
            archived = bool(item.get("archived", False))
            user.is_active = bool(item.get("is_active", True)) and not archived
            archive = db.get(StudentArchive, user.id)
            if archived and not archive:
                db.add(StudentArchive(user_id=user.id, archived_by=actor.id))
            elif not archived and archive:
                db.delete(archive)
        db.flush()

        for item in data.get("attempts") or []:
            email = str(item.get("student_email") or "").lower().strip()
            slug = str(item.get("quiz_slug") or "").strip()
            user = db.scalar(select(User).where(User.email == email, User.role == "student"))
            quiz = db.scalar(select(Quiz).where(Quiz.slug == slug).options(selectinload(Quiz.questions).selectinload(Question.choices)))
            started_at = _parse_dt(item.get("started_at"))
            if not user or not quiz or not started_at:
                continue
            attempt = db.scalar(select(QuizAttempt).where(
                QuizAttempt.user_id == user.id,
                QuizAttempt.quiz_id == quiz.id,
                QuizAttempt.started_at == started_at,
            ))
            is_new_attempt = attempt is None
            if not attempt:
                attempt = QuizAttempt(user_id=user.id, quiz_id=quiz.id, started_at=started_at)
                db.add(attempt); db.flush(); restored_attempts += 1
            attempt.submitted_at = _parse_dt(item.get("submitted_at"))
            attempt.score = item.get("score")
            attempt.level_result = item.get("level_result") or None
            attempt.needs_manual_review = bool(item.get("needs_manual_review", False))
            if is_new_attempt or not attempt.answers:
                q_by_pos = {q.position: q for q in quiz.questions}
                for a_item in item.get("answers") or []:
                    pos = a_item.get("question_position")
                    if pos is None or int(pos) not in q_by_pos:
                        continue
                    question = q_by_pos[int(pos)]
                    choice = None
                    choice_pos = a_item.get("choice_position")
                    if choice_pos is not None:
                        choice = next((c for c in question.choices if c.position == int(choice_pos)), None)
                    db.add(QuizAnswer(
                        attempt_id=attempt.id,
                        question_id=question.id,
                        choice_id=choice.id if choice else None,
                        text_answer=a_item.get("text_answer"),
                        is_correct=a_item.get("is_correct"),
                        awarded_points=a_item.get("awarded_points"),
                    ))

    _audit(db, actor.id, "safe_backup_restored", "backup", None, {
        "pages": len(data.get("pages") or []), "months": len(data.get("months") or []),
        "quizzes": len(data.get("quizzes") or []), "question_banks_restored": restored_question_banks,
        "knowledge": len(data.get("knowledge") or []),
        "students": restored_students, "attempts": restored_attempts,
    })
    db.commit()
    return {
        "ok": True,
        "note": "تمت استعادة الإعدادات والمحتوى والطلاب والنتائج المتاحة في النسخة. كلمات السر وTokens وOTP وAPI keys لا تدخل النسخة الاحتياطية حفاظًا على الأمان؛ أي طالب تم إنشاؤه من الاستعادة يقدر يعيّن كلمة سر جديدة من نسيت كلمة السر.",
        "question_banks_restored": restored_question_banks,
    }
